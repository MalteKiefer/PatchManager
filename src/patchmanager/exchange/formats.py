"""Reading and writing CSV, XLSX and JSON tables safely."""

import csv
import io
import json
import zipfile
from collections.abc import Iterable, Sequence
from typing import Any
from xml.etree.ElementTree import ParseError  # nosec B405: exception class only

from django.utils.translation import gettext as _
from openpyxl import Workbook, load_workbook

from patchmanager.accounts.services import ServiceError

MAX_ROWS = 5000
MAX_LINES = MAX_ROWS + 1000  # empty lines count too, so padding cannot bypass the limit
MAX_COLUMNS = 50
MAX_TEXT = 5_000_000  # characters in all cells: the parsed rows wait in the cache
MAX_CELL = 1000
MAX_UNZIPPED = 100 * 1024 * 1024
FORMULA_START = ("=", "+", "-", "@", "\t", "\r")  # OWASP CSV injection


def safe(value: Any) -> Any:
    """Neutralise text a spreadsheet would run as a formula."""
    if isinstance(value, str) and value.startswith(FORMULA_START):
        return "'" + value
    return value


def unsafe(value: str) -> str:
    """Undo :func:`safe` on import, so exported files import unchanged."""
    if value.startswith("'") and value[1:2] and value[1:].startswith(FORMULA_START):
        return value[1:]
    return value


def write_csv(columns: Sequence[str], rows: Iterable[dict[str, Any]]) -> bytes:
    """UTF-8 CSV with a header row."""
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    for row in rows:
        writer.writerow({k: safe("" if row.get(k) is None else row.get(k)) for k in columns})
    return out.getvalue().encode("utf-8-sig")  # BOM: spreadsheets detect UTF-8


def write_xlsx(columns: Sequence[str], rows: Iterable[dict[str, Any]], title: str) -> bytes:
    """Workbook with one sheet."""
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title[:31])
    sheet.append(list(columns))
    for row in rows:
        sheet.append([safe(row.get(k)) for k in columns])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def write_json(rows: Iterable[dict[str, Any]]) -> bytes:
    """List of objects."""
    return json.dumps(list(rows), ensure_ascii=False, indent=2, default=str).encode()


def _clean(header: Sequence[Any], values: Sequence[Any]) -> dict[str, str]:
    row = {}
    for key, value in zip(header, values, strict=False):
        if key:
            text = "" if value is None else str(value).strip()
            if len(text) > MAX_CELL:
                raise ServiceError(_("A cell is longer than %(n)s characters.") % {"n": MAX_CELL})
            row[key] = unsafe(text)
    return row


def _header(values: Sequence[Any]) -> list[str]:
    return [str(v).strip().lower() if v is not None else "" for v in values]


def read(data: bytes) -> list[dict[str, str]]:
    """Rows of the first sheet (XLSX) or the CSV, keyed by lower case header."""
    try:
        rows = _read_xlsx(data) if data[:4] == b"PK\x03\x04" else _read_csv(data)
    except (csv.Error, ParseError, IndexError) as exc:
        raise ServiceError(_("The file could not be read.")) from exc
    if sum(len(v) for row in rows for v in row.values()) > MAX_TEXT:
        raise ServiceError(_("The file contains too much text."))
    return rows


def _read_csv(data: bytes) -> list[dict[str, str]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ServiceError(_("CSV files must be UTF-8 encoded.")) from exc
    first = text.split("\n", 1)[0]
    delimiter = max(",;\t", key=first.count)  # the header has no quoted separators
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    header = _header(next(reader, []))[:MAX_COLUMNS]
    rows = []
    for number, values in enumerate(reader):
        if number > MAX_LINES:
            raise ServiceError(_("At most %(n)s rows per import.") % {"n": MAX_ROWS})
        if any(v.strip() for v in values):
            rows.append(_clean(header, values))
            if len(rows) > MAX_ROWS:
                raise ServiceError(_("At most %(n)s rows per import.") % {"n": MAX_ROWS})
    return rows


def _read_xlsx(data: bytes) -> list[dict[str, str]]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_UNZIPPED:
                raise ServiceError(_("The workbook is too large."))
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        raise ServiceError(_("The workbook could not be read.")) from exc
    try:
        sheet = workbook.worksheets[0]
        lines = sheet.iter_rows(values_only=True, max_col=MAX_COLUMNS)
        header = _header(next(lines, ()))
        rows = []
        for number, values in enumerate(lines):
            if number > MAX_LINES:
                raise ServiceError(_("At most %(n)s rows per import.") % {"n": MAX_ROWS})
            if any(v not in (None, "") for v in values):
                rows.append(_clean(header, values))
                if len(rows) > MAX_ROWS:
                    raise ServiceError(_("At most %(n)s rows per import.") % {"n": MAX_ROWS})
        return rows
    finally:
        workbook.close()
