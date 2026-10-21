import csv
import io
import json

import pytest
from django.core.cache import cache
from django.test import Client
from openpyxl import Workbook, load_workbook
from PIL import Image

from patchmanager.accounts.models import Assignment, Role
from patchmanager.accounts.services import PermissionDeniedError, ServiceError
from patchmanager.cabling import services as cabling
from patchmanager.cabling.models import Cable
from patchmanager.dcim import devices
from patchmanager.dcim.models import Device, Rack
from patchmanager.exchange import formats, importer, pdf
from patchmanager.exchange.tables import TABLES
from patchmanager.locations import services as tree
from patchmanager.locations.models import Room, Site
from patchmanager.networks.models import VLAN
from patchmanager.tenancy import rls
from tests.conftest import login

pytestmark = pytest.mark.django_db


def _csv(rows: list[dict[str, str]]) -> bytes:
    out = io.StringIO()
    writer = csv.DictWriter(out, fieldnames=list(rows[0]), delimiter=";")
    writer.writeheader()
    writer.writerows(rows)
    return out.getvalue().encode()


@pytest.fixture
def kit(make_tree, make_type, super_admin):
    t = make_tree(u_height=12)
    t["switch"] = make_type("Switch 8", u_height=1, ports=[{"name_pattern": "ge{n}", "count": 2}])
    with rls.system_context():
        sw = devices.create(
            t["request"],
            super_admin,
            device_type=t["switch"],
            room=t["room"],
            rack=t["rack"],
            position=1,
            name="sw1",
        )
        devices.create(
            t["request"],
            super_admin,
            device_type=t["switch"],
            room=t["room"],
            rack=t["rack"],
            position=2,
            name="sw2",
        )
        cabling.connect(
            t["request"],
            super_admin,
            sw.ports.get(name="ge1"),
            Device.objects.get(name="sw2").ports.get(name="ge1"),
            label="=HYPERLINK(1)",
        )
    return t


def test_formula_cells_are_neutralised_and_round_trip() -> None:
    data = formats.write_csv(["a", "b"], [{"a": "=1+1", "b": "-x"}, {"a": 5, "b": None}])
    assert data.startswith(b"\xef\xbb\xbf") and b"'=1+1" in data
    assert formats.read(data) == [{"a": "=1+1", "b": "-x"}, {"a": "5", "b": ""}]
    book = load_workbook(io.BytesIO(formats.write_xlsx(["a"], [{"a": "@SUM(A1)"}, {"a": -3}], "T")))
    assert [c.value for c in book.active["A"]] == ["a", "'@SUM(A1)", -3]
    assert formats.read(formats.write_xlsx(["a"], [{"a": "@SUM(A1)"}], "T")) == [{"a": "@SUM(A1)"}]
    assert formats.unsafe("'plain") == "'plain"


def test_read_limits_and_errors(monkeypatch) -> None:
    assert formats.read(b"name\tvid\nA\t1\n") == [{"name": "A", "vid": "1"}]
    with pytest.raises(ServiceError, match="UTF-8"):
        formats.read("name\nä".encode("latin-1"))
    monkeypatch.setattr(formats, "MAX_ROWS", 2)
    with pytest.raises(ServiceError, match="At most"):
        formats.read(b"a\n1\n2\n3\n")
    with pytest.raises(ServiceError, match="could not be read"):
        formats.read(b"a\n" + b"x" * 200_000 + b"\n")  # beyond the csv field limit
    monkeypatch.setattr(formats, "MAX_TEXT", 3)
    with pytest.raises(ServiceError, match="too much text"):
        formats.read(b"a\nabcd\n")
    monkeypatch.setattr(formats, "MAX_TEXT", 5_000_000)
    monkeypatch.setattr(formats, "MAX_LINES", 3)
    with pytest.raises(ServiceError, match="At most"):
        formats.read(b"a\n1\n\n\n\n\n\n")
    monkeypatch.setattr(formats, "MAX_LINES", 100)
    with pytest.raises(ServiceError, match="longer"):
        formats.read(b"a\n" + b"x" * 1001 + b"\n")
    with pytest.raises(ServiceError, match="could not be read"):
        formats.read(b"PK\x03\x04broken")
    book = Workbook()
    book.active.append(["a"])
    for n in range(3):
        book.active.append([n])
    out = io.BytesIO()
    book.save(out)
    with pytest.raises(ServiceError, match="At most"):
        formats.read(out.getvalue())
    monkeypatch.setattr(formats, "MAX_UNZIPPED", 10)
    with pytest.raises(ServiceError, match="too large"):
        formats.read(out.getvalue())


def test_import_locations_racks_devices_vlans_cables(kit, super_admin) -> None:
    t, req = kit, kit["request"]
    with rls.system_context():

        def run(kind, rows, apply=True):
            return [
                (r.outcome, r.message)
                for r in importer.run(req, super_admin, t["tenant"], kind, rows, apply=apply)
            ]

        loc = [
            {
                "site": "HQ",
                "building": "B2",
                "floor": "F1",
                "floor_level": "2",
                "room": "Lab",
                "room_kind": "Office",
            }
        ]
        assert run("locations", loc, apply=False) == [("created", "")]
        assert not Room.objects.filter(name="Lab").exists()  # dry run
        assert run("locations", loc) == [("created", "")]
        assert run("locations", loc) == [("unchanged", "")]
        room = Room.objects.get(name="Lab")
        assert (room.kind, room.floor.level) == ("office", 2)

        place = {"site": "hq", "building": "B2", "floor": "F1", "room": "lab"}
        assert run("racks", [{**place, "name": "L1", "u_height": "24.0", "status": "planned"}]) == [
            ("created", "")
        ]
        assert run("racks", [{**place, "name": "l1", "serial": "S9"}]) == [("updated", "")]
        assert Rack.objects.get(name="L1").serial == "S9"

        dev = {
            **place,
            "rack": "L1",
            "position": "3",
            "name": "sw3",
            "vendor": "acme",
            "model": "switch 8",
        }
        assert run("devices", [dev]) == [("created", "")]
        assert run("devices", [{**dev, "position": "5", "serial": "X1"}]) == [("updated", "")]
        sw3 = Device.objects.get(name="sw3")
        assert (sw3.space.start, sw3.serial) == (5, "X1")
        assert run("devices", [{**dev, "position": "5"}]) == [("unchanged", "")]

        assert (
            run(
                "vlans",
                [
                    {"site": "", "vid": "10", "name": "Users"},
                    {"site": "HQ", "vid": "10", "name": "Local"},
                ],
            )
            == [("created", "")] * 2
        )
        assert run("vlans", [{"site": "", "vid": "10", "name": "Staff", "status": "reserved"}]) == [
            ("updated", "")
        ]
        assert VLAN.objects.get(site=None, vid=10).name == "Staff"

        cab = {
            "device_a": "sw1",
            "port_a": "ge2",
            "device_b": "sw3",
            "port_b": "ge1",
            "media": "Cat 6",
            "length_m": "2,5",
        }
        assert run("cables", [cab]) == [("created", "")]
        assert run("cables", [{**cab, "label": "C-7"}]) == [("updated", "")]
        assert Cable.objects.get(label="C-7").length_m == 2.5


def test_import_is_all_or_nothing_and_reports_rows(kit, super_admin) -> None:
    with rls.system_context():
        rows = [
            {"site": "HQ", "building": "B1", "floor": "F1", "room": "R1", "name": "N1"},
            {"site": "HQ", "building": "B1", "floor": "F1", "room": "Nope", "name": "N2"},
            {"site": "HQ", "building": "B1", "floor": "F1", "room": "R1", "name": ""},
            {
                "site": "HQ",
                "building": "B1",
                "floor": "F1",
                "room": "R1",
                "name": "N3",
                "u_height": "x",
            },
            {
                "site": "HQ",
                "building": "B1",
                "floor": "F1",
                "room": "R1",
                "name": "N4",
                "status": "broken",
            },
        ]
        results = importer.run(
            kit["request"], super_admin, kit["tenant"], "racks", rows, apply=True
        )
        assert [r.line for r in results if r.outcome == "error"] == [3, 4, 5, 6]
        assert (
            "B1 / F1 / R1" not in results[1].message
            and "F1 / Nope not found" in results[1].message
            and "number" in results[3].message
        )
        assert "one of" in results[4].message
        assert not Rack.objects.filter(name="N1").exists()
        for bad, match in (
            ([{"site": "x", "bogus": "1"}], "Unknown"),
            ([{"building": "x"}], "Missing"),
            ([], "no rows"),
        ):
            with pytest.raises(ServiceError, match=match):
                importer.run(
                    kit["request"], super_admin, kit["tenant"], "locations", bad, apply=False
                )
        with pytest.raises(ServiceError, match="cannot be imported"):
            importer.run(
                kit["request"],
                super_admin,
                kit["tenant"],
                "ports",
                [{"device": "sw1"}],
                apply=False,
            )
        busy = {"device_a": "sw1", "port_a": "ge1", "device_b": "sw2", "port_b": "ge2"}
        assert (
            importer.run(kit["request"], super_admin, kit["tenant"], "cables", [busy], apply=False)[
                0
            ].outcome
            == "error"
        )


def test_import_respects_permissions(kit, make_user) -> None:
    user = make_user()
    with rls.system_context():
        Assignment.objects.create(
            role=Role.objects.get(key="viewer"),
            user=user,
            tenant=kit["tenant"],
            scope_path=kit["tenant"].policy_path,
        )
        with pytest.raises(PermissionDeniedError):  # viewers lack locations.site.import
            importer.run(
                kit["request"], user, kit["tenant"], "locations", [{"site": "New"}], apply=True
            )
        assert not Site.objects.filter(name="New").exists()
        assert list(TABLES["devices"].export(user, kit["tenant"], None)) == []  # nor *.export
        assert pdf  # reports are gated in the view, see test_views


def test_exports_match_import_columns_and_scope(kit, super_admin, make_user) -> None:
    with rls.system_context():
        VLAN.objects.create(tenant=kit["tenant"], vid=5, name="Mgmt")
        for kind, table in TABLES.items():
            rows = list(table.export(super_admin, kit["tenant"], None))
            assert rows, kind
            assert set(rows[0]) <= set(table.columns), kind
        devices_rows = list(TABLES["devices"].export(super_admin, kit["tenant"], kit["site"]))
        assert {r["name"] for r in devices_rows} == {"sw1", "sw2"} and devices_rows[0][
            "position"
        ] == 1
        assert (
            next(iter(TABLES["cables"].export(super_admin, kit["tenant"], None)))["label"]
            == "=HYPERLINK(1)"
        )
        # exported files import again without changes
        for kind in ("locations", "racks", "devices", "cables"):
            rows = formats.read(
                formats.write_csv(
                    TABLES[kind].columns, TABLES[kind].export(super_admin, kit["tenant"], None)
                )
            )
            results = importer.run(
                kit["request"], super_admin, kit["tenant"], kind, rows, apply=False
            )
            assert {r.outcome for r in results} == {"unchanged"}, (
                kind,
                [r.message for r in results],
            )

        user = make_user()
        secret = tree.create(kit["request"], super_admin, Room, parent=kit["floor"], name="Secret")
        devices.create(
            kit["request"], super_admin, device_type=kit["switch"], room=secret, name="hidden"
        )
        Assignment.objects.create(
            role=Role.objects.get(key="editor"),
            user=user,
            tenant=kit["tenant"],
            scope_path=kit["room"].path,
        )
        names = {r["name"] for r in TABLES["devices"].export(user, kit["tenant"], None)}
        assert names == {"sw1", "sw2"}
        hidden = Device.objects.get(name="hidden")
        cabling.connect(
            kit["request"],
            super_admin,
            Device.objects.get(name="sw1").ports.get(name="ge2"),
            hidden.ports.get(name="ge1"),
        )
        Assignment.objects.filter(user=user).update(scope_path=kit["floor"].path)
        Assignment.objects.create(
            role=Role.objects.get(key="viewer"),
            user=user,
            tenant=kit["tenant"],
            scope_path=secret.path,
            effect="deny",
        )
        ends = {
            (r["device_a"], r["device_b"])
            for r in TABLES["cables"].export(user, kit["tenant"], None)
        }
        assert ("sw1", "…") in ends and not any("hidden" in e for pair in ends for e in pair)
        # importing over an object the user cannot view reveals nothing
        results = importer.run(
            kit["request"],
            user,
            kit["tenant"],
            "devices",
            [
                {
                    "site": "HQ",
                    "building": "B1",
                    "floor": "F1",
                    "room": "R1",
                    "name": "hidden",
                    "vendor": "Acme",
                    "model": "Switch 8",
                }
            ],
            apply=False,
        )
        assert results[0].message == "hidden not found."


def test_pdf_reports(kit, super_admin, make_user, settings, tmp_path) -> None:
    logo = tmp_path / "logo.png"
    Image.new("RGB", (200, 50), "navy").save(logo)
    settings.BRAND_LOGO = str(logo)
    with rls.system_context():
        sw1 = Device.objects.get(name="sw1")
        for data in (
            pdf.site_report(super_admin, kit["site"]),
            pdf.rack_elevation(super_admin, kit["rack"]),
            pdf.port_list(super_admin, sw1),
            pdf.patch_list(super_admin, kit["site"]),
            pdf.bill_of_materials(super_admin, kit["site"]),
        ):
            assert data.startswith(b"%PDF-") and len(data) > 1000
    assert pdf.logo_problem() is None
    settings.BRAND_LOGO = str(tmp_path / "missing.png")
    assert "readable" in (pdf.logo_problem() or "")
    (tmp_path / "x.gif").write_bytes(b"GIF89a....")
    settings.BRAND_LOGO = str(tmp_path / "x.gif")
    assert "PNG or JPEG" in (pdf.logo_problem() or "")
    from patchmanager.exchange.apps import check_logo

    assert check_logo()[0].id == "patchmanager.E010"


def test_views(kit, admin_client: Client, super_admin, make_user) -> None:
    assert admin_client.get("/exchange/", secure=True).status_code == 200
    body = _csv([{"site": "HQ", "building": "B9", "floor": "", "room": ""}])
    preview = admin_client.post(
        "/exchange/import/", {"kind": "locations", "file": io.BytesIO(body)}, secure=True
    )
    assert preview.status_code == 200 and b"Import now" in preview.content
    token = preview.context["token"]
    other = Client()
    colleague = make_user()
    with rls.system_context():
        Assignment.objects.create(
            role=Role.objects.get(key="editor"),
            user=colleague,
            tenant=kit["tenant"],
            scope_path=kit["tenant"].policy_path,
        )
    login(other, colleague)
    assert (
        b"expired"
        in other.post("/exchange/import/apply/", {"token": token}, secure=True, follow=True).content
    )
    done = admin_client.post("/exchange/import/apply/", {"token": token}, secure=True, follow=True)
    assert b"1 created" in done.content
    with rls.system_context():
        assert kit["site"].buildings.filter(name="B9").exists()
    assert (
        b"expired"
        in admin_client.post(
            "/exchange/import/apply/", {"token": token}, secure=True, follow=True
        ).content
    )
    bad = admin_client.post(
        "/exchange/import/",
        {"kind": "racks", "file": io.BytesIO(b"site;name\nHQ;x\n")},
        secure=True,
    )
    assert (
        b"Missing columns" in admin_client.get("/exchange/", secure=True).content
        or bad.status_code == 302
    )

    csv_export = admin_client.get(f"/exchange/devices.csv?site={kit['site'].pk}", secure=True)
    assert csv_export["Content-Type"].startswith("text/csv") and b"sw1" in csv_export.content
    assert "attachment" in csv_export["Content-Disposition"]
    assert (
        json.loads(admin_client.get("/exchange/cables.json", secure=True).content)[0]["label"]
        == "=HYPERLINK(1)"
    )
    assert admin_client.get("/exchange/racks.xlsx", secure=True).content[:2] == b"PK"
    assert admin_client.get("/exchange/racks.exe", secure=True).status_code == 404
    assert admin_client.get("/exchange/devices.csv?site=nope", secure=True).status_code == 404
    assert admin_client.get("/exchange/templates/devices.xlsx", secure=True).content[:2] == b"PK"
    assert admin_client.get("/exchange/templates/ports.csv", secure=True).status_code == 404
    pdf_response = admin_client.get(f"/exchange/reports/rack/{kit['rack'].pk}.pdf", secure=True)
    assert pdf_response.content.startswith(b"%PDF-")
    assert (
        admin_client.get(f"/exchange/reports/nope/{kit['rack'].pk}.pdf", secure=True).status_code
        == 404
    )

    stranger = Client()
    login(stranger, make_user())
    assert (
        stranger.get(f"/exchange/reports/rack/{kit['rack'].pk}.pdf", secure=True).status_code == 404
    )


def test_import_audits_once_on_success(kit, super_admin) -> None:
    from patchmanager.audit.models import AuditEvent

    with rls.system_context():
        before = AuditEvent.objects.count()
        row = [{"site": "HQ", "building": "B7"}]
        importer.run(kit["request"], super_admin, kit["tenant"], "locations", row, apply=False)
        assert AuditEvent.objects.count() == before
        importer.run(kit["request"], super_admin, kit["tenant"], "locations", row, apply=True)
        actions = list(AuditEvent.objects.order_by("-seq").values_list("action", flat=True)[:2])
        assert actions == ["exchange.imported", "locations.building.created"]


def test_import_rate_limit(kit, admin_client: Client, monkeypatch) -> None:
    from patchmanager.exchange import views

    cache.clear()
    monkeypatch.setattr(views, "IMPORTS_PER_10_MIN", 1)
    admin_client.post(
        "/exchange/import/", {"kind": "locations", "file": io.BytesIO(b"site\nX\n")}, secure=True
    )
    response = admin_client.post(
        "/exchange/import/",
        {"kind": "locations", "file": io.BytesIO(b"site\nX\n")},
        secure=True,
        follow=True,
    )
    assert b"Too many imports" in response.content


def test_export_rate_limit(kit, admin_client: Client, monkeypatch) -> None:
    from patchmanager.exchange import views

    cache.clear()
    monkeypatch.setattr(views, "EXPORTS_PER_10_MIN", 1)
    assert admin_client.get("/exchange/racks.csv", secure=True).status_code == 200
    assert admin_client.get("/exchange/racks.csv", secure=True).status_code == 429


def test_import_never_creates_below_hidden_location(kit, make_user) -> None:
    user = make_user()
    with rls.system_context():
        Assignment.objects.create(
            role=Role.objects.get(key="editor"),
            user=user,
            tenant=kit["tenant"],
            scope_path=kit["room"].path,
        )
        row = {"site": "HQ", "building": "B1", "floor": "F1", "room": "Sneaky"}
        results = importer.run(kit["request"], user, kit["tenant"], "locations", [row], apply=False)
        assert results[0].message == "Location HQ / B1 / F1 / Sneaky not found."
        rack = {"site": "HQ", "building": "B1", "floor": "F1", "room": "R1", "name": "Own"}
        assert (
            importer.run(kit["request"], user, kit["tenant"], "racks", [rack], apply=False)[
                0
            ].outcome
            == "created"
        )


def test_viewers_cannot_export_reports(kit, make_user) -> None:
    viewer = make_user()
    with rls.system_context():
        Assignment.objects.create(
            role=Role.objects.get(key="viewer"),
            user=viewer,
            tenant=kit["tenant"],
            scope_path=kit["tenant"].policy_path,
        )
    client = Client()
    login(client, viewer)
    assert client.get(f"/exchange/reports/rack/{kit['rack'].pk}.pdf", secure=True).status_code == 404
    assert client.get(f"/exchange/reports/bom/{kit['site'].pk}.pdf", secure=True).status_code == 404
    assert b"sw1" not in client.get("/exchange/devices.csv", secure=True).content
